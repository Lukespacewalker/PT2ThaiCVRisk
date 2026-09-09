import openpyxl

excel_path = "ตรวจสุขภาพ_รปภ_ปี2569_extracted.xlsx"

wb = openpyxl.load_workbook(excel_path)
print("Sheet Names in generated Excel:")
for sheet in wb.sheetnames:
    ws = wb[sheet]
    print(f"  - Sheet: {sheet:22s} | Rows: {ws.max_row:3d} | Cols: {ws.max_column:2d}")

ws_master = wb["Master_Summary"]
print("\nSample Master Row 1 (Header):")
print([cell.value for cell in ws_master[1][:8]])

print("\nSample Master Row 2 (Employee 1):")
print([cell.value for cell in ws_master[2][:8]])
